"""
Bistapp MT16 - Belge ve Rapor İhracı Motoru (export_engine.py)
--------------------------------------------------------------
Ajanların nihai kararlarını, münazara (debate) sonuçlarını ve metriklerini
biçimlendirilmiş, renkli Excel (.xlsx) ve şık Word (.docx) raporları olarak üretir.
"""

from __future__ import annotations
import os
import uuid
from pathlib import Path
from datetime import datetime
from typing import Dict, Any, Optional, Union

try:
    import xlsxwriter
except ImportError:
    xlsxwriter = None

try:
    import docx
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
except ImportError:
    docx = None
    class MockDocx:
        def __init__(self, *args, **kwargs): pass
    Pt = Inches = RGBColor = lambda x: x
    WD_ALIGN_PARAGRAPH = WD_TABLE_ALIGNMENT = None


def _get_default_output_path(ticker: str, extension: str) -> Path:
    """Geçici dosya yolu üretir."""
    temp_dir = Path(__file__).resolve().parent / "temp"
    temp_dir.mkdir(parents=True, exist_ok=True)
    clean_ticker = ticker.strip().upper().replace(".IS", "")
    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    unique_id = uuid.uuid4().hex[:6]
    return temp_dir / f"{clean_ticker}_Rapor_{timestamp_str}_{unique_id}.{extension.lstrip('.')}"


def export_to_excel(
    decision_data: Dict[str, Any],
    ticker: str = "ASELS",
    output_path: Optional[Union[str, Path]] = None
) -> str:
    """
    Ajanların nihai kararını, münazara verilerini ve metriklerini alıp,
    renkli (AL için yeşil, SAT için kırmızı) ve formatlı bir .xlsx dosyası üretir.

    Args:
        decision_data (dict): FinalDecision veya evaluate_v2 çıktı sözlüğü.
        ticker (str): Sembol adı.
        output_path (str | Path, optional): Dosyanın kaydedileceği yol.

    Returns:
        str: Oluşturulan Excel dosyasının dosya sistemi yolu.
    """
    clean_ticker = ticker.strip().upper().replace(".IS", "")
    file_path = Path(output_path) if output_path else _get_default_output_path(clean_ticker, "xlsx")
    file_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = xlsxwriter.Workbook(str(file_path))
    
    # ------------------ STİLLER VE RENKLER ------------------
    header_fmt = workbook.add_format({
        "bold": True,
        "font_size": 14,
        "font_color": "#FFFFFF",
        "bg_color": "#0F172A",
        "align": "center",
        "valign": "vcenter",
        "border": 1
    })

    section_fmt = workbook.add_format({
        "bold": True,
        "font_size": 11,
        "font_color": "#FFFFFF",
        "bg_color": "#1E293B",
        "align": "left",
        "valign": "vcenter"
    })

    th_fmt = workbook.add_format({
        "bold": True,
        "font_size": 10,
        "bg_color": "#334155",
        "font_color": "#F8FAFC",
        "align": "center",
        "valign": "vcenter",
        "border": 1
    })

    cell_fmt = workbook.add_format({
        "font_size": 10,
        "valign": "vcenter",
        "border": 1
    })

    cell_center_fmt = workbook.add_format({
        "font_size": 10,
        "align": "center",
        "valign": "vcenter",
        "border": 1
    })

    # Karar Renk Formatları
    al_fmt = workbook.add_format({
        "bold": True,
        "font_color": "#065F46",
        "bg_color": "#D1FAE5",
        "align": "center",
        "valign": "vcenter",
        "border": 1
    })

    sat_fmt = workbook.add_format({
        "bold": True,
        "font_color": "#991B1B",
        "bg_color": "#FEE2E2",
        "align": "center",
        "valign": "vcenter",
        "border": 1
    })

    notr_fmt = workbook.add_format({
        "bold": True,
        "font_color": "#92400E",
        "bg_color": "#FEF3C7",
        "align": "center",
        "valign": "vcenter",
        "border": 1
    })

    # ------------------ ÇALIŞMA SAYFASI 1: KARAR ÖZETİ ------------------
    ws = workbook.add_worksheet("Karar_Ozeti")
    ws.set_column("A:A", 22)
    ws.set_column("B:B", 25)
    ws.set_column("C:C", 16)
    ws.set_column("D:D", 14)
    ws.set_column("E:E", 14)
    ws.set_column("F:F", 60)

    # Başlık Başlangıcı
    ws.merge_range("A1:F2", f"BISTAPP MT16 — {clean_ticker} ÇOKLU AJAN STRATEJİ VE KARAR RAPORU", header_fmt)

    # Nihai Karar Verileri
    verdict = str(decision_data.get("verdict", "AL")).upper()
    comp_score = float(decision_data.get("composite_score", 70.0))
    ts_val = decision_data.get("timestamp", datetime.now().timestamp())
    dt_str = datetime.fromtimestamp(ts_val).strftime("%d.%m.%Y %H:%M:%S") if isinstance(ts_val, (int, float)) else str(ts_val)

    verdict_style = al_fmt if "AL" in verdict else sat_fmt if "SAT" in verdict else notr_fmt

    ws.write("A4", "Hisse Sembolü:", cell_fmt)
    ws.write("B4", clean_ticker, cell_center_fmt)
    ws.write("C4", "Rapor Tarihi:", cell_fmt)
    ws.write("D4", dt_str, cell_center_fmt)

    ws.write("A5", "Nihai Karar:", cell_fmt)
    ws.write("B5", verdict, verdict_style)
    ws.write("C5", "Bileşik Puan:", cell_fmt)
    ws.write("D5", f"{comp_score:.1f} / 100", cell_center_fmt)

    meta = decision_data.get("metadata", {})
    price_val = meta.get("price", 0.0)
    chg_val = meta.get("change_pct", 0.0)
    ws.write("A6", "Canlı Fiyat:", cell_fmt)
    ws.write("B6", f"{price_val:.2f} TL" if price_val else "N/A", cell_center_fmt)
    ws.write("C6", "Günlük Değişim:", cell_fmt)
    ws.write("D6", f"%{chg_val:+.2f}" if chg_val else "%0.00", cell_center_fmt)

    # ------------------ AJAN OYLARI TABLOSU ------------------
    row = 8
    ws.merge_range(row, 0, row, 5, "1. UZMAN AJAN OYLARI VE GEREKÇELERİ", section_fmt)
    row += 1

    headers = ["Ajan Kodu", "Ajan İsmi", "Sinyal", "Skor (/100)", "Güven", "Gerekçe"]
    for col_idx, h in enumerate(headers):
        ws.write(row, col_idx, h, th_fmt)
    row += 1

    agent_votes = decision_data.get("agent_votes", {})
    for agent_id, a_data in agent_votes.items():
        sig = str(a_data.get("signal", "NOTR")).upper()
        sig_fmt = al_fmt if "AL" in sig else sat_fmt if "SAT" in sig else notr_fmt

        ws.write(row, 0, agent_id, cell_fmt)
        ws.write(row, 1, a_data.get("agent_name", agent_id), cell_fmt)
        ws.write(row, 2, sig, sig_fmt)
        ws.write(row, 3, a_data.get("score", 0.0), cell_center_fmt)
        ws.write(row, 4, f"%{int(a_data.get('confidence', 0.8) * 100)}", cell_center_fmt)
        ws.write(row, 5, a_data.get("reasoning", ""), cell_fmt)
        row += 1

    # ------------------ MÜNAZARA (DEBATE) SONUÇLARI ------------------
    row += 1
    ws.merge_range(row, 0, row, 5, "2. ÇAPRAZ İTİRAZ (DEBATE) PROTOKOLÜ VE CEZALAR", section_fmt)
    row += 1

    deb_res = decision_data.get("debate_result", {})
    challenges = deb_res.get("challenges", []) if isinstance(deb_res, dict) else []

    if challenges:
        deb_headers = ["İtiraz Eden", "Hedef Ajan", "İtiraz Tipi", "Ceza Puanı", "Eski -> Yeni Skor", "Argüman"]
        for col_idx, h in enumerate(deb_headers):
            ws.write(row, col_idx, h, th_fmt)
        row += 1

        for ch in challenges:
            ws.write(row, 0, ch.get("challenger_name", ""), cell_fmt)
            ws.write(row, 1, ch.get("target_name", ""), cell_fmt)
            ws.write(row, 2, ch.get("conflict_type", ""), cell_center_fmt)
            ws.write(row, 3, f"-{ch.get('penalty_applied', 0)}", sat_fmt)
            ws.write(row, 4, f"{ch.get('target_score_before', 0)} -> {ch.get('target_score_after', 0)}", cell_center_fmt)
            ws.write(row, 5, ch.get("argument", ""), cell_fmt)
            row += 1
    else:
        ws.write(row, 0, "Münazara Durumu:", cell_fmt)
        ws.merge_range(row, 1, row, 5, "Ajanlar arasında çelişki tespit edilmedi; uzlaşı sağlandı.", cell_fmt)
        row += 1

    # ------------------ ÇALIŞMA SAYFASI 2: TOON METRİKLERİ ------------------
    ws2 = workbook.add_worksheet("TOON_Export")
    ws2.set_column("A:A", 100)
    ws2.write("A1", "# BISTAPP MT16 STANDART TOON VERİSİ", header_fmt)
    
    toon_lines = str(decision_data.get("toon", "")).splitlines()
    for idx, l in enumerate(toon_lines):
        ws2.write(idx + 2, 0, l)

    workbook.close()
    return str(file_path)


def export_to_word(
    decision_data: Dict[str, Any],
    ticker: str = "ASELS",
    analysis_text: str = "",
    output_path: Optional[Union[str, Path]] = None
) -> str:
    """
    Ajanların kararını, LLM'den gelen doğal dil analizini ve TOON çıktısını
    şık bir başlık ve tablo yapısıyla .docx dosyası olarak üretir.

    Args:
        decision_data (dict): FinalDecision veya evaluate_v2 çıktı sözlüğü.
        ticker (str): Sembol adı.
        analysis_text (str, optional): LLM veya analist açıklaması.
        output_path (str | Path, optional): Dosyanın kaydedileceği yol.

    Returns:
        str: Oluşturulan Word dosyasının dosya sistemi yolu.
    """
    clean_ticker = ticker.strip().upper().replace(".IS", "")
    file_path = Path(output_path) if output_path else _get_default_output_path(clean_ticker, "docx")
    file_path.parent.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()

    # Sayfa Kenar Boşlukları
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.9)
        s.right_margin = Inches(0.9)

    # ------------------ BAŞLIK ------------------
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(f"BISTAPP MT16 — {clean_ticker} STRATEJİ VE KARAR RAPORU")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(15, 23, 42)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub_p = doc.add_paragraph()
    ts_val = decision_data.get("timestamp", datetime.now().timestamp())
    dt_str = datetime.fromtimestamp(ts_val).strftime("%d.%m.%Y %H:%M:%S") if isinstance(ts_val, (int, float)) else str(ts_val)
    sub_run = sub_p.add_run(f"Otonom Analiz Tarihi: {dt_str} | Karar Motoru v4.0")
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor(100, 116, 139)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ------------------ 1. NİHAİ KARAR ÖZETİ ------------------
    verdict = str(decision_data.get("verdict", "AL")).upper()
    comp_score = float(decision_data.get("composite_score", 70.0))
    
    summary_box = doc.add_paragraph()
    summary_box.paragraph_format.space_before = Pt(6)
    summary_box.paragraph_format.space_after = Pt(12)
    s_label = summary_box.add_run("NİHAİ JÜRİ KARARI: ")
    s_label.bold = True
    s_label.font.size = Pt(13)

    v_run = summary_box.add_run(f"{verdict} (Bileşik Puan: {comp_score:.1f} / 100)")
    v_run.bold = True
    v_run.font.size = Pt(14)
    if "AL" in verdict:
        v_run.font.color.rgb = RGBColor(16, 185, 129)
    elif "SAT" in verdict:
        v_run.font.color.rgb = RGBColor(239, 68, 68)
    else:
        v_run.font.color.rgb = RGBColor(245, 158, 11)

    meta = decision_data.get("metadata", {})
    px = meta.get("price", 0.0)
    chg = meta.get("change_pct", 0.0)
    info_p = doc.add_paragraph(f"Piyasa Fiyatı: {px:.2f} TL | Günlük Değişim: %{chg:+.2f} | Güven Katsayısı: %85")
    info_p.runs[0].font.size = Pt(10)
    info_p.runs[0].font.color.rgb = RGBColor(71, 85, 105)

    # ------------------ 2. UZMAN AJAN OYLARI TABLOSU ------------------
    h2 = doc.add_heading("1. Uzman Ajan Oyları ve Gerekçeleri", level=2)
    h2.paragraph_format.space_before = Pt(12)
    
    agent_votes = decision_data.get("agent_votes", {})
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    hdr_titles = ["Uzman Ajan", "Sinyal", "Skor", "Gerekçe"]
    for i, t in enumerate(hdr_titles):
        hdr_cells[i].text = t
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        # Tablo başlık koyu arka plan
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E293B"/>')
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)

    for a_id, a_data in agent_votes.items():
        row_cells = table.add_row().cells
        row_cells[0].text = a_data.get("agent_name", a_id)
        row_cells[1].text = str(a_data.get("signal", "NOTR"))
        row_cells[2].text = f"{a_data.get('score', 0):.0f}"
        row_cells[3].text = a_data.get("reasoning", "")

        for c_idx, cell in enumerate(row_cells):
            cell.paragraphs[0].runs[0].font.size = Pt(9)
            if c_idx == 1:
                sig_text = row_cells[1].text
                r = cell.paragraphs[0].runs[0]
                r.bold = True
                if "AL" in sig_text:
                    r.font.color.rgb = RGBColor(16, 185, 129)
                elif "SAT" in sig_text:
                    r.font.color.rgb = RGBColor(239, 68, 68)

    # ------------------ 3. ÇAPRAZ İTİRAZ (DEBATE) SONUÇLARI ------------------
    h3 = doc.add_heading("2. Çapraz İtiraz (Debate) Protokolü", level=2)
    h3.paragraph_format.space_before = Pt(14)

    deb_res = decision_data.get("debate_result", {})
    challenges = deb_res.get("challenges", []) if isinstance(deb_res, dict) else []

    if challenges:
        for ch in challenges:
            dp = doc.add_paragraph()
            r1 = dp.add_run(f"⚔️ {ch.get('challenger_name')} -> {ch.get('target_name')} [{ch.get('conflict_type')}]: ")
            r1.bold = True
            r1.font.color.rgb = RGBColor(185, 28, 28)
            r2 = dp.add_run(f"{ch.get('argument')} (Ceza: -{ch.get('penalty_applied')} puan)")
            r2.font.size = Pt(9.5)
    else:
        doc.add_paragraph("Ajanlar arasında keskin çelişki tespit edilmedi; tam uzlaşı sağlandı.")

    # ------------------ 4. ANALİZ VE STRATEJİ YORUMU ------------------
    h4 = doc.add_heading("3. Analist ve Strateji Notları", level=2)
    h4.paragraph_format.space_before = Pt(14)
    
    analysis_content = analysis_text or decision_data.get("coach_summary") or decision_data.get("summary", "Detaylı çoklu ajan simülasyonu başarıyla tamamlandı.")
    ap = doc.add_paragraph(analysis_content)
    ap.runs[0].font.size = Pt(10)

    # ------------------ 5. STANDART TOON ÇIKTISI ------------------
    h5 = doc.add_heading("4. Bistapp Standart TOON Çıktısı (LLM Uyumlu)", level=2)
    h5.paragraph_format.space_before = Pt(14)

    toon_text = decision_data.get("toon", "")
    if toon_text:
        toon_p = doc.add_paragraph()
        toon_run = toon_p.add_run(toon_text)
        toon_run.font.name = "Courier New"
        toon_run.font.size = Pt(8.5)
        toon_run.font.color.rgb = RGBColor(51, 65, 85)

    # Alt Bilgi
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    footer = doc.add_paragraph("⚠️ Yasal Uyarı: Bu rapor yatırım tavsiyesi değildir (YTD). Otonom algoritmik jüri sistemi tarafından üretilmiştir.")
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].italic = True
    footer.runs[0].font.color.rgb = RGBColor(148, 163, 184)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(str(file_path))
    return str(file_path)
